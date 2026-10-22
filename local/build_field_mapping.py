"""Build a reviewer-friendly Topia -> Equus field mapping workbook from the
migration's mapping report (field names only, no employee data)."""
import collections
import glob
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "lambda" / "src"))
import mapping_config as cfg  # noqa: E402

report = openpyxl.load_workbook(glob.glob(str(REPO / "local_s3/dest_s3/equus/reports/*.xlsx"))[-1])
cm = list(report["Column Mapping"].iter_rows(values_only=True))[1:]
cov = list(report["Equus Coverage"].iter_rows(values_only=True))[1:]
rule_by_header = {r[1]: r[2] for r in cov}
topia_by_letter = {r[0]: r[3] or r[2] or "" for r in cm}

STATUS = {  # status -> (label, fill colour, explanation)
    "mapped": ("Mapped", "C6EFCE", "Copied to the Equus column shown"),
    "address": ("Mapped - address", "DDEBF7", "Becomes an UPDATE row (mailing address)"),
    "contact": ("Mapped - contact", "DDEBF7", "Becomes an UPDATE row (employee contact)"),
    "dependent": ("Mapped - dependent", "DDEBF7", "Becomes an UPDATE row (dependent)"),
    "rule": ("Rule", "FFF2CC", "Drives a derived value"),
    "ignored": ("Excluded", "EDEDED", "Deliberately not migrated"),
    "duplicate": ("Duplicate", "F8CBAD", "Target already filled by another column"),
    "unmapped": ("Not mapped", "FFC7CE", "No Equus field identified"),
}


def how(status, note):
    note = note or ""
    if status != "mapped":
        return note
    if note == "row 2 = row 3":
        return "Direct: Equus field named in Topia row 2"
    if note.startswith("override"):
        return "Override in mapping_config.FIELD_OVERRIDES"
    if note.startswith("alias of"):
        return f"Renamed: {note.removeprefix('alias of ')} -> Equus spelling"
    return note


HDR_FONT = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(wb, title, headers, rows, widths, status_col=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HDR_FONT, HDR_FILL, Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = WRAP
        if status_col is not None:
            colour = next((v[1] for v in STATUS.values() if v[0] == r[status_col].value), None)
            if colour:
                r[status_col].fill = PatternFill("solid", fgColor=colour)
    ws.freeze_panes = "B2" if title != "Summary" else None
    ws.auto_filter.ref = ws.dimensions
    return ws


def add_review(ws, decision_col, n):
    dv = DataValidation(type="list", formula1='"Approve,Change,Remove,Question"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{decision_col}2:{decision_col}{n + 1}")


wb = openpyxl.Workbook()
wb.remove(wb.active)

# --- Topia -> Equus ------------------------------------------------------
t_rows = []
for letter, key, ui, topia, status, target, tcol, note in cm:
    label = STATUS.get(status, (status,))[0]
    rule = rule_by_header.get(target) if target else None
    t_rows.append([letter, topia or ui or "", ui or "", key or "", label,
                   target or "", tcol or "", rule or "", how(status, note), "", ""])
ws_t = sheet(wb, "Topia to Equus",
             ["Topia Col", "Topia Field (row 4)", "Equus UI Name (Topia row 3)",
              "Equus Field per Topia row 2 / note", "Status", "Mapped To Equus Header",
              "Equus Col", "Equus Conversion Rule (row 1)", "How Mapped / Reason",
              "Reviewer Decision", "Reviewer Comments"],
             t_rows, [9, 34, 30, 40, 16, 40, 9, 34, 42, 16, 40], status_col=4)
add_review(ws_t, "J", len(t_rows))

# --- Equus fields --------------------------------------------------------
CONTROL = {cfg.COL_EMPLOYEE_ACTION, cfg.COL_ASSIGNMENT_COMMAND, cfg.COL_ASSIGNMENT_NUMBER,
           cfg.COL_ASSIGNMENT_IDENTIFIER}
e_rows = []
for col, header, rule, populated, sources in cov:
    letters = [s.strip() for s in str(sources or "").split(",") if s.strip()]
    if header in CONTROL:
        status = "Control column"
        src = "Set by the migration (INSERT / NEW / Topia Assignment ID)"
    elif letters:
        status = "Mapped"
        src = "; ".join(f"{l}: {topia_by_letter.get(l, '')}" for l in letters)
    else:
        status = "Not mapped"
        src = ""
    e_rows.append([col, header, rule or "", status, src, "", ""])
ws_e = sheet(wb, "Equus Fields",
             ["Equus Col", "Equus Header (row 3)", "Conversion Rule (row 1)", "Status",
              "Topia Source (col: field)", "Reviewer Decision", "Reviewer Comments"],
             e_rows, [9, 42, 40, 16, 50, 16, 40], status_col=3)
for r in ws_e.iter_rows(min_row=2):
    if r[3].value == "Control column":
        r[3].fill = PatternFill("solid", fgColor="DDEBF7")
add_review(ws_e, "F", len(e_rows))

# --- Value rules ---------------------------------------------------------
v_rows = []
for header, m in cfg.VALUE_MAPS.items():
    for src, dst in m.items():
        v_rows.append(["Dropdown value", header, src, dst])
for label, loc in cfg.ADDRESS_COLUMNS.items():
    v_rows.append(["Address -> LOCATION_TYPE", "ASSIGNMENT_MAILING_ADDRESS", label, loc])
for label, (loc, typ) in cfg.CONTACT_COLUMNS.items():
    v_rows.append(["Contact -> LOCATION/CONTACT TYPE", "ASSIGNMENT_EMPLOYEE_CONTACT", label, f"{loc} / {typ}"])
v_rows.append(["Derived rule", "ASSIGNMENT!PHASE", "Cancellation Reason is not empty", cfg.CANCELLED_PHASE])
v_rows.append(["Blank values", "All fields", ", ".join(sorted(cfg.BLANK_TOKENS)), "(empty)"])
v_rows.append(["Country codes", "Match Country Code columns", "Country name", "ISO 3166-1 alpha-3 (e.g. United States -> USA)"])
v_rows.append(["Dates", "Date columns", "Any date", "US format per Equus row 1 rule"])
v_rows.append(["Resources", "Coordinator / additional resources",
               "Topia name or email", "Equus username - RESOURCE_MAP is empty, needs a list" if not cfg.RESOURCE_MAP else "see RESOURCE_MAP"])
sheet(wb, "Value Rules", ["Rule Type", "Equus Field", "Topia Value", "Equus Value"], v_rows, [30, 40, 50, 50])

# --- Summary -------------------------------------------------------------
tc = collections.Counter(r[4] for r in t_rows)
ec = collections.Counter(r[3] for r in e_rows)
s = [["Topia -> Equus field mapping (for review)", ""], ["", ""],
     ["Topia columns", len(t_rows)]]
s += [[f"  {STATUS[k][0]}", tc[STATUS[k][0]]] for k in STATUS if tc[STATUS[k][0]]]
s += [["", ""], ["Equus columns", len(e_rows)]] + [[f"  {k}", v] for k, v in ec.items()]
s += [["", ""], ["Status legend", ""]] + [[v[0], v[2]] for v in STATUS.values()]
s += [["", ""], ["How to review",
       "Use the 'Reviewer Decision' dropdown (Approve / Change / Remove / Question) and add "
       "comments on the 'Topia to Equus' and 'Equus Fields' sheets. Mapping changes go into "
       "lambda/src/mapping_config.py."]]
ws_s = sheet(wb, "Summary", ["Item", "Value"], s, [34, 90])
ws_s.auto_filter.ref = None
ws_s["A2"].font = Font(bold=True, size=14)
for r in ws_s.iter_rows(min_row=2):
    if r[0].value in [v[0] for v in STATUS.values()]:
        r[0].fill = PatternFill("solid", fgColor=next(v[1] for v in STATUS.values() if v[0] == r[0].value))
wb.move_sheet("Summary", offset=-3)
wb.active = 0

out = REPO / "Topia_Equus_Field_Mapping.xlsx"
wb.save(out)
print(out)
for row in s[2:16]:
    print("  ", row)
