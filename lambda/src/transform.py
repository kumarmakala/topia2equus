"""Transform a Topia Full Historical Past Assignment export into the Equus
EAFH import layout. Pure functions: bytes in, bytes out, no AWS calls."""

from __future__ import annotations

import io
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

import mapping_config as cfg

KEY_RE = re.compile(r"^[A-Z0-9_]+![A-Z0-9_]+$")
DATE_FORMAT = "MM/DD/YYYY"
_DATE_PATTERNS = ("%m/%d/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d %b %Y", "%Y-%m-%d %H:%M:%S", "%m/%d/%y")

ADDR = "ASSIGNMENT_MAILING_ADDRESS!"
CONTACT = "ASSIGNMENT_EMPLOYEE_CONTACT!"


def norm_label(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value)).strip().lower() if value is not None else ""


def norm_key(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    key = value.strip().upper()
    if not KEY_RE.match(key):
        return None
    for pattern, repl in cfg.KEY_ALIASES:
        key = re.sub(pattern, repl, key)
    return key


def is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == "" or value.strip().lower() in cfg.BLANK_TOKENS
    return False


@dataclass
class SourceColumn:
    idx: int
    key: Any
    ui: Any
    topia: Any
    status: str = "unmapped"  # mapped | address | contact | dependent | rule | ignored | unmapped | duplicate
    target: str | None = None
    note: str = ""

    @property
    def letter(self) -> str:
        return get_column_letter(self.idx)

    @property
    def label(self) -> str:
        return norm_label(self.topia) or norm_label(self.ui)


@dataclass
class TargetLayout:
    headers: dict[str, int]  # header -> column index
    rules: dict[str, str]  # header -> lower-cased row-1 instruction

    def is_date(self, header: str) -> bool:
        return "date format" in self.rules.get(header, "") or header.endswith("_DATE")

    def is_country(self, header: str) -> bool:
        return "country code" in self.rules.get(header, "")

    def is_resource(self, header: str) -> bool:
        return "RESOURCE" in header

    def is_dropdown(self, header: str) -> bool:
        return "dropdown" in self.rules.get(header, "") or "yesno" in self.rules.get(header, "")


@dataclass
class Report:
    columns: list[SourceColumn] = field(default_factory=list)
    warnings: list[tuple[Any, str, str]] = field(default_factory=list)
    dropdown_values: Counter = field(default_factory=Counter)
    coverage: Counter = field(default_factory=Counter)
    assignments: int = 0
    rows_written: int = 0

    def warn(self, row: Any, fld: str, msg: str) -> None:
        self.warnings.append((row, fld, msg))


# --------------------------------------------------------------------------- layout

def read_target_layout(ws) -> TargetLayout:
    headers, rules = {}, {}
    for c in range(1, ws.max_column + 1):
        raw = ws.cell(cfg.TARGET_HEADER_ROW, c).value
        if raw is None:
            continue
        header = norm_key(raw) or str(raw).strip()
        if header in headers:
            raise ValueError(f"Duplicate Equus header {header!r} in template")
        headers[header] = c
        rules[header] = norm_label(ws.cell(cfg.TARGET_RULE_ROW, c).value)
    return TargetLayout(headers, rules)


def build_column_plan(ws, layout: TargetLayout) -> list[SourceColumn]:
    cols = []
    for c in range(cfg.SOURCE_FIRST_COL, ws.max_column + 1):
        col = SourceColumn(
            idx=c,
            key=ws.cell(cfg.SOURCE_KEY_ROW, c).value,
            ui=ws.cell(cfg.SOURCE_UI_ROW, c).value,
            topia=ws.cell(cfg.SOURCE_TOPIA_ROW, c).value,
        )
        if col.key is None and col.ui is None and col.topia is None:
            continue
        cols.append(col)

    claimed: dict[str, SourceColumn] = {}

    def claim(col: SourceColumn, target: str, note: str) -> None:
        if target not in layout.headers:
            col.status, col.note = "unmapped", f"override target {target} not in Equus template"
        elif target in claimed:
            col.status, col.target = "duplicate", target
            col.note = f"{target} already filled from source column {claimed[target].letter}"
        else:
            col.status, col.target, col.note = "mapped", target, note
            claimed[target] = col

    # Pass 1: special roles and explicit overrides (these win over row 2).
    for col in cols:
        label = col.label
        if label in cfg.ADDRESS_COLUMNS:
            col.status, col.note = "address", f"LOCATION_TYPE = {cfg.ADDRESS_COLUMNS[label]}"
        elif label in cfg.CONTACT_COLUMNS:
            col.status, col.note = "contact", "LOCATION_TYPE/CONTACT_TYPE = %s / %s" % cfg.CONTACT_COLUMNS[label]
        elif label.startswith(cfg.IGNORED_LABEL_PREFIXES):
            col.status, col.note = "ignored", "describes a constant location type"
        elif label == cfg.CANCELLATION_REASON_LABEL:
            col.status, col.note = "rule", f"non-empty -> ASSIGNMENT!PHASE = {cfg.CANCELLED_PHASE}"
        elif label in cfg.FIELD_OVERRIDES or norm_label(col.ui) in cfg.FIELD_OVERRIDES:
            target = cfg.FIELD_OVERRIDES.get(label, cfg.FIELD_OVERRIDES.get(norm_label(col.ui)))
            if target is None:
                col.status, col.note = "ignored", "excluded in mapping_config.FIELD_OVERRIDES"
            else:
                claim(col, target, "override (row 2 not usable)")

    # Pass 2: row-2 technical field vs Equus row-3 header.
    for col in cols:
        if col.status != "unmapped":
            continue
        key = norm_key(col.key)
        if key is None:
            col.note = "row 2 has no technical field" if col.key is None else "row 2 is a note, not a field"
        elif key.startswith("ASSIGNMENT_DEPENDENT"):
            col.status, col.target = "dependent", key
            if key not in layout.headers:
                col.status, col.note = "unmapped", f"{key} not in Equus template"
        elif key in layout.headers:
            claim(col, key, "row 2 = row 3" if key == str(col.key).strip().upper() else f"alias of {col.key}")
        else:
            col.note = f"{key} not in Equus template"
    return cols


# --------------------------------------------------------------------------- values

def to_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date() if value.time() == datetime.min.time() else value
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):  # Excel serial
        return openpyxl.utils.datetime.from_excel(value)
    text = str(value).strip()
    for fmt in _DATE_PATTERNS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def convert(header: str, value: Any, layout: TargetLayout, report: Report, row: Any) -> Any:
    if is_blank(value):
        return None
    if isinstance(value, str):
        value = value.strip()

    if layout.is_date(header):
        parsed = to_date(value)
        if parsed is None:
            report.warn(row, header, f"unparseable date {value!r}; left blank")
        return parsed

    if layout.is_country(header):
        text = str(value)
        if text.upper() in cfg.ISO3_CODES:
            return text.upper()
        iso = cfg.COUNTRY_ISO3.get(text.lower())
        if iso is None:
            report.warn(row, header, f"no ISO-3 code for country {text!r}; passed through")
            return text
        return iso

    if isinstance(value, bool):
        value = "Yes" if value else "No"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if "CF_TEXT" in header and not isinstance(value, str):
        value = str(value)

    if header in cfg.VALUE_MAPS and isinstance(value, str):
        value = cfg.VALUE_MAPS[header].get(value.lower(), value)
    if layout.is_resource(header) and isinstance(value, str):
        value = cfg.RESOURCE_MAP.get(value, value)
    if layout.is_dropdown(header) or layout.is_resource(header):
        report.dropdown_values[(header, value)] += 1
    return value


# --------------------------------------------------------------------------- rows

def _group_dependents(cols: list[SourceColumn]) -> list[list[SourceColumn]]:
    """Split repeated ASSIGNMENT_DEPENDENT* columns into one group per child."""
    groups: list[list[SourceColumn]] = []
    for col in cols:
        if col.status != "dependent":
            continue
        if not groups or any(g.target == col.target for g in groups[-1]):
            groups.append([])
        groups[-1].append(col)
    return groups


def build_records(src_ws, cols: list[SourceColumn], layout: TargetLayout, report: Report) -> list[dict]:
    mapped = [c for c in cols if c.status == "mapped"]
    addresses = [c for c in cols if c.status == "address"]
    contacts = [c for c in cols if c.status == "contact"]
    dependents = _group_dependents(cols)
    cancel_col = next((c for c in cols if c.status == "rule"), None)
    ident_col = next((c for c in cols if c.label == cfg.ASSIGNMENT_IDENTIFIER_LABEL), None)

    out_rows: list[dict] = []
    seen_employees: set = set()

    for r in range(cfg.SOURCE_DATA_START_ROW, src_ws.max_row + 1):
        raw = {c.idx: src_ws.cell(r, c.idx).value for c in cols}
        if all(is_blank(v) for v in raw.values()):
            continue
        row_ref = f"source row {r}"
        main = {}
        for col in mapped:
            main[col.target] = convert(col.target, raw[col.idx], layout, report, row_ref)

        employee = main.get(cfg.COL_EMPLOYEE_NUMBER)
        identifier = convert("ASSIGNMENT_CF!CF_TEXT029", raw[ident_col.idx], layout, report, row_ref) if ident_col else None
        if identifier is None:
            identifier = employee
            report.warn(row_ref, cfg.COL_ASSIGNMENT_IDENTIFIER, "no Topia Assignment ID; using Employee Number")
        if employee is None:
            report.warn(row_ref, cfg.COL_EMPLOYEE_NUMBER, "missing Employee Number (Talent Link ID)")

        # Business rules
        if cancel_col and not is_blank(raw[cancel_col.idx]):
            main["ASSIGNMENT!PHASE"] = cfg.CANCELLED_PHASE
        home, host = main.get("ASSIGNMENT!FROM_COUNTRY"), main.get("ASSIGNMENT!TO_COUNTRY")
        if main.get("ASSIGNMENT!INTERNATIONAL_RELO") is None and home and host:
            main["ASSIGNMENT!INTERNATIONAL_RELO"] = "Yes" if home != host else "No"
        start = main.get("ASSIGNMENT!START_DATE")
        if cfg.DEFAULT_DATES_FROM_START and start:
            if main.get("ASSIGNMENT_COMPENSATION_SALARY!ANNUAL_BASE_PAY_AMOUNT") is not None:
                main["ASSIGNMENT_COMPENSATION_SALARY!EFFECTIVE_DATE"] = main.get("ASSIGNMENT_COMPENSATION_SALARY!EFFECTIVE_DATE") or start
            if any(main.get(k) is not None for k in ("ASSIGNMENT_FAMILY_SIZE!HOST_FAMILY_SIZE", "ASSIGNMENT_FAMILY_SIZE!HOME_FAMILY_SIZE")):
                main["ASSIGNMENT_FAMILY_SIZE!EFFECTIVE_DATE"] = main.get("ASSIGNMENT_FAMILY_SIZE!EFFECTIVE_DATE") or start
        begin = start if cfg.DEFAULT_DATES_FROM_START else None

        # Child records
        addr_recs = []
        for col in addresses:
            value = raw[col.idx]
            if is_blank(value):
                continue
            lines = [ln.strip() for ln in str(value).splitlines() if ln.strip()]
            rec = {
                ADDR + "PRIMARY_ADDRESS": "Yes",
                ADDR + "LOCATION_TYPE": cfg.ADDRESS_COLUMNS[col.label],
                ADDR + "BEGIN_DATE": begin,
                ADDR + "LINE_1": lines[0] if lines else None,
                ADDR + "LINE_2": lines[1] if len(lines) > 1 else None,
                ADDR + "LINE_3": ", ".join(lines[2:]) or None,
            }
            report.warn(row_ref, ADDR + "LINE_1", f"address from column {col.letter} loaded as free text; verify city/postal/country split")
            addr_recs.append(rec)

        contact_recs = []
        for col in contacts:
            value = raw[col.idx]
            if is_blank(value):
                continue
            loc, ctype = cfg.CONTACT_COLUMNS[col.label]
            contact_recs.append({
                CONTACT + "LOCATION_TYPE": loc,
                CONTACT + "CONTACT_TYPE": ctype,
                CONTACT + "NUMBER_ADDRESS": str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip(),
                CONTACT + "START_DATE": begin,
                CONTACT + "PRIMARY_CONTACT": "Yes",
            })

        dep_recs = []
        for group in dependents:
            rec = {c.target: convert(c.target, raw[c.idx], layout, report, row_ref) for c in group}
            if any(v is not None for v in rec.values()):
                dep_recs.append(rec)

        # Equus layout: first row INSERTs the assignment; extra child records
        # go on UPDATE rows keyed by the Assignment Identifier.
        n_rows = max(1, len(addr_recs), len(contact_recs), len(dep_recs))
        for i in range(n_rows):
            row = dict(main) if i == 0 else {}
            if i == 0:
                row[cfg.COL_EMPLOYEE_ACTION] = cfg.EMPLOYEE_ACTION_FIRST if employee not in seen_employees else cfg.EMPLOYEE_ACTION_REPEAT
                row[cfg.COL_ASSIGNMENT_COMMAND] = "INSERT"
                row[cfg.COL_ASSIGNMENT_NUMBER] = cfg.NEW_ASSIGNMENT_NUMBER
            else:
                row[cfg.COL_EMPLOYEE_ACTION] = cfg.EMPLOYEE_ACTION_REPEAT
                row[cfg.COL_ASSIGNMENT_COMMAND] = "UPDATE"
            row[cfg.COL_ASSIGNMENT_IDENTIFIER] = identifier
            for recs in (addr_recs, contact_recs, dep_recs):
                if i < len(recs):
                    row.update(recs[i])
            out_rows.append(row)
        seen_employees.add(employee)
        report.assignments += 1
    return out_rows


# --------------------------------------------------------------------------- io

def write_output(template_bytes: bytes, rows: list[dict], layout_ws_name: str, report: Report) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    ws = wb[layout_ws_name]
    layout = read_target_layout(ws)
    if ws.max_row >= cfg.TARGET_DATA_START_ROW:
        ws.delete_rows(cfg.TARGET_DATA_START_ROW, ws.max_row - cfg.TARGET_DATA_START_ROW + 1)
    for offset, row in enumerate(rows):
        r = cfg.TARGET_DATA_START_ROW + offset
        for header, value in row.items():
            if value is None:
                continue
            c = layout.headers.get(header)
            if c is None:
                report.warn(f"output row {r}", header, "no Equus column for value; dropped")
                continue
            cell = ws.cell(r, c, value)
            if isinstance(value, (date, datetime)):
                cell.number_format = DATE_FORMAT
            report.coverage[header] += 1
    report.rows_written = len(rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_report(report: Report, layout: TargetLayout) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Column Mapping"
    ws.append(["Source Col", "Row 2 (Equus technical)", "Row 3 (Equus UI)", "Row 4 (Topia)", "Status", "Equus Header", "Equus Col", "Note"])
    for c in report.columns:
        target_col = get_column_letter(layout.headers[c.target]) if c.target in layout.headers else None
        ws.append([c.letter, c.key, c.ui, c.topia, c.status, c.target, target_col, c.note])

    ws = wb.create_sheet("Equus Coverage")
    ws.append(["Equus Col", "Equus Header", "Row 1 Rule", "Cells Populated", "Source Columns"])
    sources = defaultdict(list)
    for c in report.columns:
        if c.target:
            sources[c.target].append(c.letter)
    for header, idx in sorted(layout.headers.items(), key=lambda kv: kv[1]):
        ws.append([get_column_letter(idx), header, layout.rules.get(header), report.coverage.get(header, 0), ", ".join(sources.get(header, []))])

    ws = wb.create_sheet("Dropdown Values")
    ws.append(["Equus Header", "Value Sent", "Occurrences"])
    for (header, value), n in sorted(report.dropdown_values.items(), key=lambda kv: (kv[0][0], str(kv[0][1]))):
        ws.append([header, value, n])

    ws = wb.create_sheet("Warnings")
    ws.append(["Where", "Field", "Message"])
    for w in report.warnings:
        ws.append(list(w))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def transform(source_bytes: bytes, template_bytes: bytes) -> tuple[bytes, bytes, Report]:
    """Return (equus_workbook_bytes, report_workbook_bytes, report)."""
    src_wb = openpyxl.load_workbook(io.BytesIO(source_bytes), data_only=True)
    if cfg.SOURCE_SHEET not in src_wb.sheetnames:
        raise ValueError(f"Source sheet {cfg.SOURCE_SHEET!r} not found; got {src_wb.sheetnames}")
    src_ws = src_wb[cfg.SOURCE_SHEET]

    tpl_wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    layout = read_target_layout(tpl_wb[cfg.TARGET_SHEET])

    report = Report()
    report.columns = build_column_plan(src_ws, layout)
    rows = build_records(src_ws, report.columns, layout, report)
    out = write_output(template_bytes, rows, cfg.TARGET_SHEET, report)
    return out, build_report(report, layout), report
