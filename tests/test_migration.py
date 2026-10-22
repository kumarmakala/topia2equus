import io
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "lambda" / "src"))

import handler  # noqa: E402
from transform import transform  # noqa: E402

SOURCE = ROOT / "tests" / "fixtures" / "topia_sample.xlsx"
TEMPLATE = ROOT / "lambda" / "src" / "templates" / "equus_eafh_template.xlsx"


def _rows(xlsx: bytes):
    ws = openpyxl.load_workbook(io.BytesIO(xlsx))["Sheet1"]
    headers = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(4, ws.max_row + 1):
        rows.append({h: ws.cell(r, c).value for c, h in enumerate(headers, 1) if ws.cell(r, c).value is not None})
    return ws, rows


@pytest.fixture(scope="module")
def result():
    return transform(SOURCE.read_bytes(), TEMPLATE.read_bytes())


def test_template_headers_preserved_and_sample_rows_removed(result):
    out, _, report = result
    ws, rows = _rows(out)
    tpl = openpyxl.load_workbook(TEMPLATE)["Sheet1"]
    for r in (1, 2, 3):
        assert [ws.cell(r, c).value for c in range(1, ws.max_column + 1)] == \
               [tpl.cell(r, c).value for c in range(1, tpl.max_column + 1)]
    assert report.assignments == 2
    assert len(rows) == 2


def test_core_fields(result):
    _, rows = _rows(result[0])
    raj = rows[0]
    assert raj["A"] == "INSERT" and raj["Assignment Command"] == "INSERT"
    assert raj["Assignment Number"] == "NEW"
    assert raj["Employee Number"] == 10001
    assert raj["Assignment Identifier"] == "90001"
    assert raj["ASSIGNMENT_CF!CF_TEXT029"] == "90001"
    assert raj["EMPLOYEE!BIRTH_DATE"].date() == date(1980, 1, 15)


def test_conversions_and_rules(result):
    out, _, _ = result
    ws, rows = _rows(out)
    raj, nic = rows
    assert (raj["ASSIGNMENT!FROM_COUNTRY"], raj["ASSIGNMENT!TO_COUNTRY"]) == ("EGY", "IND")
    assert nic["EMPLOYEE!PRIMARY_CITIZENSHIP"] == "USA"
    assert raj["EMPLOYEE!EMPLOYMENT_STATUS"] == "Active"
    assert raj["EMPLOYEE!TYPE"] == "Employee"
    assert raj["ASSIGNMENT!PHASE"] == "Cancelled"
    assert raj["ASSIGNMENT!INTERNATIONAL_RELO"] == "Yes"
    assert raj["ASSIGNMENT!FROM_COMPANY_SEGMENT_LEVEL_1"] == "SCHNEIDER ELECTRIC EGYPT"
    assert raj["ASSIGNMENT!GROUP_NAME"] == "Nova Project Denmark"
    # Host Job Level is tagged TO_TITLE in row 2; override keeps title intact.
    assert nic["ASSIGNMENT!TO_TITLE"] == "Solution Architect"
    assert nic["ASSIGNMENT!TO_SALARY_GRADE"] == 8
    assert raj["EMPLOYEE_CF!CF_TEXT014"] == "900000000000001"
    # "Not Provided" placeholders are blanked.
    assert "EMPLOYEE_CF!CF_TEXT005" not in raj
    date_col = [c for c in range(1, ws.max_column + 1) if ws.cell(3, c).value == "ASSIGNMENT!START_DATE"][0]
    assert ws.cell(4, date_col).number_format == "MM/DD/YYYY"


def test_report_lists_unmapped_columns(result):
    _, report_bytes, report = result
    wb = openpyxl.load_workbook(io.BytesIO(report_bytes))
    assert wb.sheetnames == ["Column Mapping", "Equus Coverage", "Dropdown Values", "Warnings"]
    unmapped = {c.topia for c in report.columns if c.status == "unmapped"}
    assert "Host Annualized Base Pay Amount" in unmapped  # OFFER_SALARY not in Equus file


def _source_with_children() -> bytes:
    wb = openpyxl.load_workbook(SOURCE)
    ws = wb["Full Historical Past Assignment"]
    labels = {ws.cell(4, c).value: c for c in range(1, ws.max_column + 1)}
    ws.cell(5, labels["Personal Home Address with Effective date"], "12 Big Street\nDown the Road")
    ws.cell(5, labels["Work Host Address with Effective date"], "1 Office Park")
    ws.cell(5, labels["Mobile Phone Number"], 1234567.0)
    # Dependent columns appended the way the "Dependent Import" sheet lays them out.
    start = ws.max_column + 1
    for i, (key, label, value) in enumerate([
        ("ASSIGNMENT_DEPENDENT!FIRST_NAME", "Child 1 First Name", "Child1"),
        ("ASSIGNMENT_DEPENDENT!BIRTH_DATE", "Child 1 Date of Birth", datetime(2010, 1, 1)),
        ("ASSIGNMENT_DEPENDENT!FIRST_NAME", "Child 2 First Name", "Child2"),
        ("ASSIGNMENT_DEPENDENT!BIRTH_DATE", "Child 2 Date of Birth", datetime(2012, 1, 1)),
    ]):
        ws.cell(2, start + i, key)
        ws.cell(4, start + i, label)
        ws.cell(5, start + i, value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_child_records_become_update_rows():
    out, _, _ = transform(_source_with_children(), TEMPLATE.read_bytes())
    _, rows = _rows(out)
    raj = [r for r in rows if r.get("Assignment Identifier") == "90001"]
    assert len(raj) == 2
    first, second = raj
    assert first["ASSIGNMENT_MAILING_ADDRESS!LOCATION_TYPE"] == "Departure Residence"
    assert first["ASSIGNMENT_MAILING_ADDRESS!LINE_2"] == "Down the Road"
    assert first["ASSIGNMENT_EMPLOYEE_CONTACT!NUMBER_ADDRESS"] == "1234567"
    assert first["ASSIGNMENT_DEPENDENT!FIRST_NAME"] == "Child1"
    assert second["A"] == "IGNORE" and second["Assignment Command"] == "UPDATE"
    assert "Employee Number" not in second
    assert second["ASSIGNMENT_MAILING_ADDRESS!LOCATION_TYPE"] == "Destination Office"
    assert second["ASSIGNMENT_DEPENDENT!FIRST_NAME"] == "Child2"


# --------------------------------------------------------------------------- handler

class FakeS3:
    def __init__(self, objects):
        self.objects = objects
        self.puts = {}

    def get_object(self, Bucket, Key):
        return {"Body": io.BytesIO(self.objects[(Bucket, Key)])}

    def put_object(self, Bucket, Key, Body, **kwargs):
        self.puts[(Bucket, Key)] = Body


def _event(bucket, key, size=1000):
    return {"Records": [{"eventName": "ObjectCreated:Put",
                         "s3": {"bucket": {"name": bucket}, "object": {"key": key, "size": size}}}]}


@pytest.fixture
def fake_s3(monkeypatch):
    s3 = FakeS3({("src", "topia/incoming/FULL HISTORICAL.xlsx"): SOURCE.read_bytes()})
    monkeypatch.setattr(handler, "_s3", s3)
    monkeypatch.setenv("DEST_BUCKET", "dst")
    monkeypatch.setenv("DEST_PREFIX", "equus/")
    monkeypatch.setenv("SOURCE_PREFIX", "topia/incoming/")
    monkeypatch.delenv("TEMPLATE_S3_URI", raising=False)
    return s3


def test_handler_s3_event_writes_output_and_report(fake_s3):
    res = handler.lambda_handler(_event("src", "topia/incoming/FULL+HISTORICAL.xlsx"), None)
    summary = res["results"][0]
    assert summary["assignments"] == 2
    keys = sorted(k for _, k in fake_s3.puts)
    assert keys[0].startswith("equus/FULL HISTORICAL_EQUUS_")
    assert keys[1].startswith("equus/reports/FULL HISTORICAL_MAPPING_REPORT_")


@pytest.mark.parametrize("key,reason", [
    ("other/FULL HISTORICAL.xlsx", "prefix"),
    ("topia/incoming/notes.csv", "end with"),
    ("topia/incoming/~$FULL HISTORICAL.xlsx", "lock file"),
])
def test_handler_filter_skips(fake_s3, key, reason):
    res = handler.lambda_handler(_event("src", key), None)
    assert reason in res["results"][0]["skipped"]
    assert fake_s3.puts == {}


def test_handler_skips_oversized_and_own_outputs(fake_s3, monkeypatch):
    res = handler.lambda_handler(_event("src", "topia/incoming/big.xlsx", size=10**9), None)
    assert "larger" in res["results"][0]["skipped"]
    monkeypatch.setenv("SOURCE_PREFIX", "")
    res = handler.lambda_handler(_event("dst", "equus/x_EQUUS_1.xlsx"), None)
    assert "own outputs" in res["results"][0]["skipped"]
